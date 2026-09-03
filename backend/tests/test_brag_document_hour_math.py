"""Table-driven checks of aggregate_member_month's deterministic hour arithmetic — the
numbers this feature's trust story depends on (see models/brag_document_job.py's
docstring): they are computed once here and never touched by an LLM anywhere downstream.
"""

import io

from openpyxl import Workbook

from app.services.standup_excel_service import aggregate_member_month, parse_standup_workbook

_DAY_COLUMNS = [(3, 4, 5, 6), (7, 8, 9, 10), (11, 12, 13, 14), (15, 16, 17, 18), (19, 20, 21, 22)]


def _build_single_week_workbook(sheet_title: str, headers: list[str], done_texts: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    for (done_col, *_rest), value in zip(_DAY_COLUMNS, headers):
        ws.cell(row=1, column=done_col, value=value)
    ws.cell(row=3, column=1, value="Jane Doe")
    ws.cell(row=3, column=2, value="")
    for (done_col, *_rest), text in zip(_DAY_COLUMNS, done_texts):
        ws.cell(row=3, column=done_col, value=text)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_full_week_no_holiday_no_leave():
    data = _build_single_week_workbook(
        "Week 32 (Aug 3 - Aug 7)",
        ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"],
        [f"Billable: X\n- Task {i} 8hrs" for i in range(5)],
    )
    workbook = parse_standup_workbook(data, "standup.xlsx")
    summary = aggregate_member_month(workbook, "Jane Doe", "August 2026")

    assert summary.total_hours == 40.0
    assert summary.gross_base_hours == 40.0  # 5 gross days * 8h
    assert summary.holiday_deducted_hours == 0.0
    assert summary.expected_target_hours == 40.0
    assert summary.balance_hours == 0.0
    assert summary.target_completion_pct == 100.0


def test_one_holiday_reduces_expected_target_and_deducts_hours():
    data = _build_single_week_workbook(
        "Week 32 (Aug 3 - Aug 7)",
        ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"],
        [
            "Company Holiday",
            "Billable: X\n- Task 8hrs",
            "Billable: X\n- Task 8hrs",
            "Billable: X\n- Task 8hrs",
            "Billable: X\n- Task 8hrs",
        ],
    )
    workbook = parse_standup_workbook(data, "standup.xlsx")
    summary = aggregate_member_month(workbook, "Jane Doe", "August 2026")

    assert summary.total_hours == 32.0
    assert summary.gross_base_hours == 40.0  # still 5 gross days in the month
    assert summary.holiday_count == 1
    assert summary.holiday_deducted_hours == 8.0
    assert summary.expected_target_hours == 32.0  # 40 - 8
    assert summary.balance_hours == 0.0  # exactly meets the reduced target
    assert summary.target_completion_pct == 100.0


def test_leave_day_does_not_reduce_expected_target_but_shows_as_shortfall():
    """A personal leave day, unlike a company holiday, does NOT reduce the expected
    target — it shows up as a negative balance instead."""
    data = _build_single_week_workbook(
        "Week 32 (Aug 3 - Aug 7)",
        ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"],
        [
            "On Leave",
            "Billable: X\n- Task 8hrs",
            "Billable: X\n- Task 8hrs",
            "Billable: X\n- Task 8hrs",
            "Billable: X\n- Task 8hrs",
        ],
    )
    workbook = parse_standup_workbook(data, "standup.xlsx")
    summary = aggregate_member_month(workbook, "Jane Doe", "August 2026")

    assert summary.leave_count == 1
    assert summary.total_hours == 32.0
    assert summary.expected_target_hours == 40.0  # unaffected by leave
    assert summary.balance_hours == -8.0
    assert summary.target_completion_pct == 80.0


def test_billable_and_non_billable_hours_split():
    data = _build_single_week_workbook(
        "Week 32 (Aug 3 - Aug 7)",
        ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"],
        ["Billable: X\n- Task 6hrs", "Non-Billable: Internal\n- Meeting 2hrs", "", "", ""],
    )
    workbook = parse_standup_workbook(data, "standup.xlsx")
    summary = aggregate_member_month(workbook, "Jane Doe", "August 2026")

    assert summary.billable_hours == 6.0
    assert summary.non_billable_hours == 2.0
    assert summary.total_hours == 8.0


def test_blocker_is_counted():
    # Set the blocker cell (block_col of day 1 = column 5) directly — helper functions
    # above only fill the Done column, not Plan/Blocker/Delivery.
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 32 (Aug 3 - Aug 7)"
    for (done_col, *_rest), value in zip(_DAY_COLUMNS, ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"]):
        ws.cell(row=1, column=done_col, value=value)
    ws.cell(row=3, column=1, value="Jane Doe")
    ws.cell(row=3, column=2, value="")
    ws.cell(row=3, column=3, value="Billable: X\n- Task 1hr")
    ws.cell(row=3, column=5, value="Blocked on API access")
    buf = io.BytesIO()
    wb.save(buf)

    workbook = parse_standup_workbook(buf.getvalue(), "standup.xlsx")
    summary = aggregate_member_month(workbook, "Jane Doe", "August 2026")
    assert summary.blocker_count == 1


def test_month_with_no_matching_weeks_has_zero_everything():
    data = _build_single_week_workbook(
        "Week 32 (Aug 3 - Aug 7)",
        ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"],
        ["Billable: X\n- Task 8hrs", "", "", "", ""],
    )
    workbook = parse_standup_workbook(data, "standup.xlsx")
    summary = aggregate_member_month(workbook, "Jane Doe", "December 2026")

    assert summary.total_hours == 0.0
    assert summary.expected_target_hours == 0.0
    assert summary.target_completion_pct == 0.0  # 0 target, 0 logged -> defined as 0%, not 100%


def test_unknown_member_returns_zeroed_summary_not_an_error():
    data = _build_single_week_workbook(
        "Week 32 (Aug 3 - Aug 7)",
        ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"],
        ["Billable: X\n- Task 8hrs", "", "", "", ""],
    )
    workbook = parse_standup_workbook(data, "standup.xlsx")
    summary = aggregate_member_month(workbook, "Nobody Here", "August 2026")

    assert summary.total_hours == 0.0
    assert summary.included_weeks == []

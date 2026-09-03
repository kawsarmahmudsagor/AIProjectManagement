"""Pure unit tests for standup_excel_service — no DB/Postgres needed. Builds minimal
.xlsx workbooks in-memory via openpyxl.Workbook() and feeds their bytes through
parse_standup_workbook(), mirroring the reference standup_cli tool's own scenarios but
against the openpyxl-based reimplementation (see standup_excel_service.py's docstring for
why header parsing must handle both string and native datetime/date cell values).
"""

import io
from datetime import datetime

from openpyxl import Workbook

from app.services.standup_excel_service import (
    aggregate_member_month,
    best_member_match,
    normalize_month_name,
    parse_day_date_header,
    parse_standup_workbook,
    parse_week_info,
)

_DAY_COLUMNS = [(3, 4, 5, 6), (7, 8, 9, 10), (11, 12, 13, 14), (15, 16, 17, 18), (19, 20, 21, 22)]


def _set_headers(ws, headers: list) -> None:
    """headers: 5 values (str or datetime), one Done-column date header per weekday."""
    for (done_col, *_rest), value in zip(_DAY_COLUMNS, headers):
        ws.cell(row=1, column=done_col, value=value)


def _set_member_row(ws, row: int, name: str, pool_tasks: str, day_texts: list[tuple[str, str, str, str]]) -> None:
    ws.cell(row=row, column=1, value=name)
    ws.cell(row=row, column=2, value=pool_tasks)
    for (done_col, plan_col, block_col, deliv_col), (done, plan, block, deliv) in zip(_DAY_COLUMNS, day_texts):
        ws.cell(row=row, column=done_col, value=done)
        ws.cell(row=row, column=plan_col, value=plan)
        ws.cell(row=row, column=block_col, value=block)
        ws.cell(row=row, column=deliv_col, value=deliv)


def _to_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- week-name / date-header parsing -------------------------------------------------


def test_parse_week_info_extracts_number_and_keeps_label():
    num, label = parse_week_info("Week 32 (Aug 3 - Aug 7)")
    assert num == 32
    assert label == "Week 32 (Aug 3 - Aug 7)"


def test_normalize_month_name_handles_abbreviations_and_unknowns():
    assert normalize_month_name("Aug 3") == "August"
    assert normalize_month_name("September") == "September"
    assert normalize_month_name("nonsense") is None
    assert normalize_month_name("") is None


def test_parse_day_date_header_handles_string_header():
    date_str, month, year = parse_day_date_header("Aug 3", "Week 32 (Aug 3 - Aug 7) 2026")
    assert date_str == "Aug 3"
    assert month == "August"
    assert year == 2026


def test_parse_day_date_header_handles_native_datetime_header():
    """openpyxl returns a native datetime.datetime for a header cell Excel formatted as a
    date, where the reference tool's raw-XML parser only ever saw a plain string — the
    exact case this module's docstring calls out."""
    header_value = datetime(2026, 8, 3)  # noqa: DTZ001 — openpyxl itself returns naive datetimes
    date_str, month, year = parse_day_date_header(header_value, "Week 32 (Aug 3 - Aug 7)")
    assert month == "August"
    assert year == 2026
    assert "3" in date_str


def test_parse_day_date_header_falls_back_to_sheet_name_for_bare_number():
    _date_str, month, _year = parse_day_date_header("3", "Week 32 (Aug 3 - Aug 7)")
    assert month == "August"


def test_parse_day_date_header_handles_blank_header():
    date_str, month, _year = parse_day_date_header(None, "Week 32 (Aug 3 - Aug 7)")
    assert date_str == "Unknown Date"
    assert month == "August"


# --- full workbook parse ---------------------------------------------------------------


def test_parse_standup_workbook_finds_members_and_months():
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 32 (Aug 3 - Aug 7)"
    _set_headers(ws, ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"])
    _set_member_row(
        ws,
        3,
        "Jane Doe",
        "",
        [
            ("Billable:\n- XR23: Fixed shader bug 4hrs", "", "", ""),
            ("- Wrote unit tests 3hrs", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
        ],
    )
    workbook = parse_standup_workbook(_to_bytes(wb), "standup.xlsx")

    assert "Jane Doe" in workbook.team_members
    assert "August 2026" in workbook.available_months
    assert workbook.sheet_names == ["Week 32 (Aug 3 - Aug 7)"]


def test_ignores_total_row():
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 32 (Aug 3 - Aug 7)"
    _set_headers(ws, ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"])
    # Row 3+ is real data; a literal "Total" summary row must be excluded from team_members.
    _set_member_row(ws, 3, "Total", "", [("", "", "", "")] * 5)
    _set_member_row(ws, 4, "Jane Doe", "", [("- Did a thing 1hr", "", "", "")] + [("", "", "", "")] * 4)

    workbook = parse_standup_workbook(_to_bytes(wb), "standup.xlsx")
    assert workbook.team_members == ["Jane Doe"]


def test_billable_category_carries_forward_to_subsequent_bullets():
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 32 (Aug 3 - Aug 7)"
    _set_headers(ws, ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"])
    _set_member_row(
        ws,
        3,
        "Jane Doe",
        "",
        [
            ("Billable: XR23\n- Fixed shader bug 4hrs\n- Wrote unit tests 3hrs", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
        ],
    )
    workbook = parse_standup_workbook(_to_bytes(wb), "standup.xlsx")
    day1_tasks = workbook.members_data["Jane Doe"][0].days[0].tasks

    assert len(day1_tasks) == 2
    assert all(t.is_billable for t in day1_tasks)
    assert day1_tasks[0].hours == 4.0
    assert day1_tasks[1].hours == 3.0
    # Domain-guessing (_infer_project_domain) is deliberately NOT ported — the category
    # is carried forward verbatim from the category line, never re-classified.
    assert day1_tasks[0].category == "Billable"


def test_non_billable_category_is_detected():
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 32 (Aug 3 - Aug 7)"
    _set_headers(ws, ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"])
    _set_member_row(
        ws,
        3,
        "Jane Doe",
        "",
        [
            ("Non-Billable: Internal\n- Team meeting 1hr", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
        ],
    )
    workbook = parse_standup_workbook(_to_bytes(wb), "standup.xlsx")
    task = workbook.members_data["Jane Doe"][0].days[0].tasks[0]
    assert task.is_billable is False
    assert task.hours == 1.0


def test_holiday_day_is_detected_and_excludes_tasks():
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 32 (Aug 3 - Aug 7)"
    _set_headers(ws, ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"])
    _set_member_row(
        ws,
        3,
        "Jane Doe",
        "",
        [
            ("Company Holiday", "", "", ""),
            ("- Fixed a bug 2hrs", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
        ],
    )
    workbook = parse_standup_workbook(_to_bytes(wb), "standup.xlsx")
    day1 = workbook.members_data["Jane Doe"][0].days[0]
    assert day1.is_holiday is True
    assert day1.tasks == []


def test_leave_day_is_detected_and_excludes_tasks():
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 32 (Aug 3 - Aug 7)"
    _set_headers(ws, ["Aug 3", "Aug 4", "Aug 5", "Aug 6", "Aug 7"])
    _set_member_row(
        ws,
        3,
        "Jane Doe",
        "",
        [
            ("On Leave", "", "", ""),
            ("- Fixed a bug 2hrs", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
            ("", "", "", ""),
        ],
    )
    workbook = parse_standup_workbook(_to_bytes(wb), "standup.xlsx")
    day1 = workbook.members_data["Jane Doe"][0].days[0]
    assert day1.is_leave is True
    assert day1.tasks == []


def test_month_boundary_week_is_recognized_and_split_correctly():
    wb = Workbook()
    ws = wb.active
    ws.title = "Week 32 (Jul 30 - Aug 3)"
    _set_headers(ws, ["Jul 30", "Jul 31", "Aug 1", "Aug 2", "Aug 3"])
    _set_member_row(
        ws,
        3,
        "Jane Doe",
        "",
        [
            (f"Billable: X\n- Task {i} 8hrs", "", "", "")
            for i in range(5)
        ],
    )
    workbook = parse_standup_workbook(_to_bytes(wb), "standup.xlsx")
    assert "July 2026" in workbook.available_months
    assert "August 2026" in workbook.available_months

    july_summary = aggregate_member_month(workbook, "Jane Doe", "July 2026")
    august_summary = aggregate_member_month(workbook, "Jane Doe", "August 2026")

    assert july_summary.total_hours == 16.0  # Jul 30 + Jul 31
    assert august_summary.total_hours == 24.0  # Aug 1, 2, 3


# --- fuzzy member match ----------------------------------------------------------------


def test_best_member_match_exact():
    name, ratio = best_member_match(["Jane Doe", "John Smith"], "Jane Doe")
    assert name == "Jane Doe"
    assert ratio == 1.0


def test_best_member_match_close_variant_scores_above_threshold():
    name, ratio = best_member_match(["Jane A. Doe", "John Smith"], "Jane Doe")
    assert name == "Jane A. Doe"
    assert ratio > 0.72


def test_best_member_match_no_close_match_scores_below_threshold():
    _name, ratio = best_member_match(["Completely Different Person"], "Jane Doe")
    assert ratio < 0.72


def test_best_member_match_empty_inputs():
    assert best_member_match([], "Jane Doe") == (None, 0.0)
    assert best_member_match(["Jane Doe"], "") == (None, 0.0)

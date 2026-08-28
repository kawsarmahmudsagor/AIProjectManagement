"""DOCX export via python-docx + html-for-docx.

Naming trap (docs/RESEARCH.md §C4): the maintained package installs as `html-for-docx`
on PyPI but its import name is `html4docx` — the similarly-named `htmldocx` and
`html4docx` packages *on PyPI* are both abandoned forks. Don't let an editor
auto-import fix swap this.
"""

import io

from docx import Document
from html4docx import HtmlToDocx

from app.models.project import Project


def _add_section(document: Document, parser: HtmlToDocx, heading: str, html: str) -> None:
    if not html.strip():
        return
    document.add_heading(heading, level=2)
    parser.add_html_to_document(html, document)


def render_project_docx(project: Project) -> bytes:
    document = Document()
    parser = HtmlToDocx()

    title = document.add_heading(project.name, level=1)
    title.add_run(f" — {project.role}").italic = True

    date_range = (
        f"{project.start_date:%B %Y} – "
        f"{'Present' if project.is_current else (project.end_date and f'{project.end_date:%B %Y}') or ''}"
    )
    document.add_paragraph(date_range)

    if project.technologies:
        document.add_paragraph("Technologies: " + ", ".join(project.technologies))

    if project.project_url:
        document.add_paragraph(project.project_url)

    _add_section(document, parser, "Project Description", project.description_long_html or f"<p>{project.description_short_html}</p>")
    _add_section(document, parser, "Responsibilities", project.responsibilities_long_html or f"<p>{project.responsibilities_short_html}</p>")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()

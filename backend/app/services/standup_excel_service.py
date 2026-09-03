"""Parses a team-wide weekly standup Excel workbook and aggregates one member's month.

Ported from the reference desktop tool (`standup_cli`)'s `parsers/data_models.py`,
`parsers/xlsx_parser.py`, and `parsers/date_detector.py`, reimplemented against
`openpyxl` instead of raw `zipfile`/`xml.etree.ElementTree`. Two deliberate departures
from the reference:

1. `_infer_project_domain`/`KNOWN_PROJECT_NAMES` (a hardcoded keyword classifier tuned to
   one specific team) is dropped entirely — this app replaces that guesswork with an LLM
   call grounded in the user's own saved `Project.responsibilities_*` text (see
   `brag_document_service.py`). Only the deterministic billable/non-billable category-line
   carry-forward is kept.
2. Header-date cells: openpyxl may return a native `datetime.datetime`/`datetime.date`
   object for a column header Excel formatted as a date, where the reference's raw-XML
   parser only ever saw the cell's literal string content (`<v>` text). Every function that
   reads a header cell here accepts either shape.

`aggregate_member_month`'s hour arithmetic (gross/expected/holiday-deducted/balance/
completion-%) is ported as close to verbatim as the data model allows and is never touched
by an LLM anywhere in this feature — see `models/brag_document_job.py`'s docstring for why
`hour_stats` is a separate column from the LLM-authored `result`.
"""

from __future__ import annotations

import difflib
import io
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

# --- Data models (ported from standup_cli/parsers/data_models.py) -------------------


@dataclass
class StandupTask:
    category: str
    task_name: str
    hours: float
    raw_text: str
    is_billable: bool = False


@dataclass
class DayEntry:
    date_str: str  # e.g., 'Aug 3'
    full_date_label: str  # e.g., 'Week 32 (Aug 3 - Aug 7) | Aug 3'
    month_name: str  # e.g., 'August'
    year: int
    done_text: str = ""
    plan_text: str = ""
    blocker_text: str = ""
    delivery_text: str = ""
    is_leave: bool = False
    is_holiday: bool = False
    holiday_name: str = ""
    tasks: list[StandupTask] = field(default_factory=list)

    @property
    def total_hours(self) -> float:
        return sum(t.hours for t in self.tasks)

    @property
    def billable_hours(self) -> float:
        return sum(t.hours for t in self.tasks if t.is_billable)

    @property
    def non_billable_hours(self) -> float:
        return sum(t.hours for t in self.tasks if not t.is_billable)


@dataclass
class MemberWeekData:
    member_name: str
    sheet_name: str
    week_number: int
    weekly_pool_tasks: str
    days: list[DayEntry] = field(default_factory=list)


@dataclass
class WeekSummaryStat:
    sheet_name: str
    logged_hours: float = 0.0
    target_hours: float = 40.0
    gross_days_in_month: int = 5
    holiday_days: int = 0
    holiday_names: list[str] = field(default_factory=list)
    leave_days: int = 0
    leave_dates: list[str] = field(default_factory=list)
    balance_hours: float = 0.0
    completion_percentage: float = 0.0


@dataclass
class MemberMonthlySummary:
    member_name: str
    month_name: str
    year: int
    included_weeks: list[str]
    days: list[DayEntry] = field(default_factory=list)
    project_hours: dict[str, float] = field(default_factory=dict)
    weekly_hours: dict[str, float] = field(default_factory=dict)
    weekly_stats: dict[str, WeekSummaryStat] = field(default_factory=dict)

    total_hours: float = 0.0
    expected_target_hours: float = 0.0
    gross_base_hours: float = 0.0
    holiday_deducted_hours: float = 0.0
    holiday_count: int = 0
    holiday_names: list[str] = field(default_factory=list)
    leave_count: int = 0
    leave_hours: float = 0.0
    blocker_count: int = 0
    leave_dates: list[str] = field(default_factory=list)
    blockers: list[tuple[str, str]] = field(default_factory=list)
    billable_hours: float = 0.0
    non_billable_hours: float = 0.0

    @property
    def balance_hours(self) -> float:
        return self.total_hours - self.expected_target_hours

    @property
    def target_completion_pct(self) -> float:
        if self.expected_target_hours <= 0:
            return 100.0 if self.total_hours > 0 else 0.0
        return (self.total_hours / self.expected_target_hours) * 100.0


@dataclass
class StandupWorkbook:
    file_path: str
    sheet_names: list[str]
    team_members: list[str]
    available_months: list[str]  # e.g. ['July 2026', 'August 2026']
    members_data: dict[str, list[MemberWeekData]] = field(default_factory=dict)
    company_holidays: dict[tuple[str, str, str], str] = field(default_factory=dict)


# --- Date/week helpers (ported from standup_cli/parsers/date_detector.py) -----------

MONTH_NAMES_FULL = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]

MONTH_ABBR_MAP = {
    "jan": "January", "feb": "February", "mar": "March", "apr": "April",
    "may": "May", "jun": "June", "jul": "July", "aug": "August",
    "sep": "September", "sept": "September", "oct": "October",
    "nov": "November", "dec": "December",
}

_MONTH_RE = re.compile(
    r"\b(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|"
    r"Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\b",
    re.IGNORECASE,
)
_YEAR_RE = re.compile(r"\b(20\d\d)\b")
_WEEK_RE = re.compile(r"Week\s*(\d+)", re.IGNORECASE)

DEFAULT_YEAR = 2026


def normalize_month_name(text: str | None) -> str | None:
    """Finds a month name from a string like 'Aug', 'August', 'Jul 27'."""
    if not text:
        return None
    match = _MONTH_RE.search(text)
    if match:
        abbr = match.group(1)[:3].lower()
        return MONTH_ABBR_MAP.get(abbr)
    return None


def extract_year_from_filename_or_text(text: str, fallback_year: int = DEFAULT_YEAR) -> int:
    """Extracts a 4-digit year like 2026 from a filename or sheet title."""
    match = _YEAR_RE.search(text or "")
    if match:
        return int(match.group(1))
    return fallback_year


def parse_week_info(sheet_name: str) -> tuple[int, str]:
    """Extracts the week number and label from a sheet name like 'Week 32 (Aug 3 - Aug 7)'."""
    week_match = _WEEK_RE.search(sheet_name)
    week_num = int(week_match.group(1)) if week_match else 0
    return week_num, sheet_name


def parse_day_date_header(
    date_header: object, sheet_name: str, default_year: int = DEFAULT_YEAR
) -> tuple[str, str, int]:
    """Parses a column date-header cell value into (cleaned_date_str, month_name, year).

    `date_header` may be a plain string (e.g. 'Aug 3', '31') — the only shape the
    reference's raw-XML parser ever saw — or a native `datetime.datetime`/`datetime.date`
    object, which openpyxl returns whenever Excel has that cell formatted as a date. Both
    are handled here; a typed date is strictly more reliable (it carries its own real
    year) than falling back to the sheet name/filename the way the string path must.
    """
    if isinstance(date_header, datetime):
        date_header = date_header.date()
    if isinstance(date_header, date):
        month_name = MONTH_NAMES_FULL[date_header.month - 1]
        date_str = f"{month_name[:3]} {date_header.day}"
        return date_str, month_name, date_header.year

    if date_header is None or not str(date_header).strip():
        sheet_month = normalize_month_name(sheet_name) or "Unknown Month"
        return "Unknown Date", sheet_month, default_year

    date_str = str(date_header).strip()
    month = normalize_month_name(date_str)
    if not month:
        # A bare number like "3" or "31" — infer the month from the sheet name instead.
        month = normalize_month_name(sheet_name) or "August"

    year = extract_year_from_filename_or_text(sheet_name, default_year)
    return date_str, month, year


# --- Task-string parsing (ported from standup_cli/parsers/xlsx_parser.py) -----------
#
# _infer_project_domain/KNOWN_PROJECT_NAMES are deliberately NOT ported — see this
# module's docstring. Only the billable/non-billable category-line carry-forward
# survives; the assigned category is whatever category line preceded the bullet
# (verbatim), never a guessed project domain.

_HOUR_RE = re.compile(r"[-–—:]*\s*(\d+(?:\.\d+)?)\s*h(?:ou)?r?s?", re.IGNORECASE)
_HOUR_STRIP_RE = re.compile(r"[-–—:]*\s*\d+(?:\.\d+)?\s*h(?:ou)?r?s?.*$", re.IGNORECASE)
_BULLET_STRIP_RE = re.compile(r"^[-*•\s]+")


def _parse_task_string(text: str) -> list[StandupTask]:
    """Parses a standup cell's raw text into discrete tasks with category, hours, and
    billable flag."""
    if not text or not text.strip():
        return []

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    tasks: list[StandupTask] = []
    current_category = "General"
    is_billable = False

    for line in lines:
        if not line.startswith("-") and not line.startswith("*") and not line.startswith("•"):
            cat_match = line.split(":")[0].strip()
            if cat_match:
                current_category = cat_match
                is_billable = "non-billable" not in cat_match.lower()
            continue

        hour_match = _HOUR_RE.search(line)
        hours = float(hour_match.group(1)) if hour_match else 0.0

        task_name = _HOUR_STRIP_RE.sub("", line)
        task_name = _BULLET_STRIP_RE.sub("", task_name).strip()
        if not task_name:
            task_name = line

        tasks.append(
            StandupTask(
                category=current_category,
                task_name=task_name,
                hours=hours,
                raw_text=line,
                is_billable=is_billable,
            )
        )

    return tasks


_HOLIDAY_KEYWORDS = ("holiday", "uprising day", "eid e miladunnabi", "eid")
_LEAVE_TEXT_VALUES = {""}
_BLANK_BLOCKER_VALUES = {"none", "no", "n/a", "-", ""}

# 1-based column indices for the 5 weekday blocks: (Done, Plan, Blocker, Delivery),
# mirroring the reference's column-letter groups C-F, G-J, K-N, O-R, S-V.
_DAY_COLUMN_GROUPS = [(3, 4, 5, 6), (7, 8, 9, 10), (11, 12, 13, 14), (15, 16, 17, 18), (19, 20, 21, 22)]

_IGNORED_NAME_VALUES = {"team member", "name", "total"}


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def parse_standup_workbook(data: bytes, filename: str) -> StandupWorkbook:
    """Parses raw `.xlsx` bytes into a `StandupWorkbook`. Uses openpyxl in read-only mode
    (cheap for a workbook with many weekly sheets); rows are read via `iter_rows()` into a
    plain dict-per-row structure rather than random `cell()` access, which is the well-
    supported pattern for openpyxl's read-only worksheets."""
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        filename_year = extract_year_from_filename_or_text(filename, DEFAULT_YEAR)

        members_set: set[str] = set()
        members_data: dict[str, list[MemberWeekData]] = defaultdict(list)
        all_months_found: set[str] = set()
        company_holidays: dict[tuple[str, str, str], str] = {}

        sheet_names = list(workbook.sheetnames)

        for sheet_name in sheet_names:
            ws = workbook[sheet_name]
            week_num, _week_label = parse_week_info(sheet_name)

            # enumerate()'d indices, not cell.row/cell.column: a read-only worksheet fills
            # gaps in its used range with lightweight `EmptyCell` placeholders that have no
            # `.row`/`.column` attribute of their own (unlike a normal writable-mode Cell).
            # min_row/min_col=1 keeps row/column numbering stable even if the sheet's used
            # range doesn't start at A1 for some reason.
            rows: dict[int, dict[int, object]] = {}
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0
            if max_row and max_col:
                for r_idx, row in enumerate(
                    ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col), start=1
                ):
                    rows[r_idx] = {c_idx: cell.value for c_idx, cell in enumerate(row, start=1)}

            if not rows:
                continue

            row1 = rows.get(1, {})

            for r_idx, row_cols in rows.items():
                if r_idx < 3:
                    continue

                member_name = _cell_text(row_cols.get(1))
                if not member_name or member_name.lower() in _IGNORED_NAME_VALUES:
                    continue

                members_set.add(member_name)
                pool_tasks = _cell_text(row_cols.get(2))
                day_entries: list[DayEntry] = []

                for done_col, plan_col, block_col, deliv_col in _DAY_COLUMN_GROUPS:
                    date_hdr = row1.get(done_col)
                    date_str, month_name, year = parse_day_date_header(date_hdr, sheet_name, filename_year)

                    if month_name:
                        all_months_found.add(f"{month_name} {year}")

                    done_val = _cell_text(row_cols.get(done_col))
                    plan_val = _cell_text(row_cols.get(plan_col))
                    block_val = _cell_text(row_cols.get(block_col))
                    deliv_val = _cell_text(row_cols.get(deliv_col))

                    if block_val.lower() in _BLANK_BLOCKER_VALUES:
                        block_val = ""

                    is_leave = False
                    is_holiday = False
                    holiday_name = ""
                    combined_text = f"{done_val} {plan_val}".lower()

                    if any(k in combined_text for k in _HOLIDAY_KEYWORDS):
                        is_holiday = True
                        if "uprising day" in combined_text:
                            holiday_name = "July Uprising Day"
                        elif "eid" in combined_text:
                            holiday_name = "Eid E Miladunnabi"
                        else:
                            holiday_name = "Company Holiday"
                        company_holidays[(sheet_name, date_str, month_name)] = holiday_name
                    elif "on leave" in combined_text or combined_text.strip() == "leave" or "leave" in combined_text.split():
                        is_leave = True

                    tasks = [] if (is_leave or is_holiday) else _parse_task_string(done_val)
                    full_label = f"{sheet_name} | {date_str}"

                    day_entries.append(
                        DayEntry(
                            date_str=date_str,
                            full_date_label=full_label,
                            month_name=month_name,
                            year=year,
                            done_text=done_val,
                            plan_text=plan_val,
                            blocker_text=block_val,
                            delivery_text=deliv_val,
                            is_leave=is_leave,
                            is_holiday=is_holiday,
                            holiday_name=holiday_name,
                            tasks=tasks,
                        )
                    )

                members_data[member_name].append(
                    MemberWeekData(
                        member_name=member_name,
                        sheet_name=sheet_name,
                        week_number=week_num,
                        weekly_pool_tasks=pool_tasks,
                        days=day_entries,
                    )
                )

        month_order = {m: i for i, m in enumerate(MONTH_NAMES_FULL)}

        def _sort_month_key(m_str: str) -> tuple[int, int]:
            parts = m_str.split()
            m_name = parts[0]
            year = int(parts[1]) if len(parts) > 1 else DEFAULT_YEAR
            return (year, month_order.get(m_name, 99))

        return StandupWorkbook(
            file_path=filename,
            sheet_names=sheet_names,
            team_members=sorted(members_set),
            available_months=sorted(all_months_found, key=_sort_month_key),
            members_data=dict(members_data),
            company_holidays=company_holidays,
        )
    finally:
        workbook.close()


def aggregate_member_month(
    workbook: StandupWorkbook, member_name: str, target_month_year: str
) -> MemberMonthlySummary:
    """Aggregates one member's daily records/hours for a selected month (e.g. 'August
    2026'). Ported verbatim from the reference tool's arithmetic — standard 40h/week base
    (8h/day), minus 8h per recognized company holiday; a personal leave day is tracked but
    does NOT reduce the expected target (it shows up as a shortfall instead). This
    function is pure Python and is never delegated to or overwritten by an LLM."""
    target_parts = target_month_year.split()
    target_month = target_parts[0]
    target_year = int(target_parts[1]) if len(target_parts) > 1 else DEFAULT_YEAR

    weeks_data = workbook.members_data.get(member_name, [])

    summary = MemberMonthlySummary(
        member_name=member_name, month_name=target_month, year=target_year, included_weeks=[]
    )

    included_sheets_order: list[str] = []
    seen_holidays: set[str] = set()

    for w in weeks_data:
        matching_days = [
            day
            for day in w.days
            if day.month_name.lower() == target_month.lower() and (day.year == target_year or day.year == 0)
        ]
        if not matching_days:
            continue

        if w.sheet_name not in included_sheets_order:
            included_sheets_order.append(w.sheet_name)

        w_stat = WeekSummaryStat(sheet_name=w.sheet_name, gross_days_in_month=len(matching_days), target_hours=0.0)

        for day in matching_days:
            summary.days.append(day)

            holiday_key = (w.sheet_name, day.date_str, day.month_name)
            is_company_holiday = day.is_holiday or (holiday_key in workbook.company_holidays)

            if is_company_holiday:
                w_stat.holiday_days += 1
                h_name = day.holiday_name or workbook.company_holidays.get(holiday_key, "Company Holiday")
                h_full = f"{h_name} ({day.date_str})"
                if h_full not in seen_holidays:
                    seen_holidays.add(h_full)
                    summary.holiday_names.append(h_full)
                if h_name not in w_stat.holiday_names:
                    w_stat.holiday_names.append(h_name)
            elif day.is_leave:
                summary.leave_count += 1
                summary.leave_dates.append(day.full_date_label)
                w_stat.leave_days += 1
                w_stat.leave_dates.append(day.date_str)

            if day.blocker_text:
                summary.blocker_count += 1
                summary.blockers.append((day.full_date_label, day.blocker_text))

            for t in day.tasks:
                summary.total_hours += t.hours
                w_stat.logged_hours += t.hours
                if t.is_billable:
                    summary.billable_hours += t.hours
                else:
                    summary.non_billable_hours += t.hours

                cat = t.category.strip()
                summary.project_hours[cat] = summary.project_hours.get(cat, 0.0) + t.hours

        gross_week_hours = w_stat.gross_days_in_month * 8.0
        holiday_week_deduction = w_stat.holiday_days * 8.0
        w_stat.target_hours = max(0.0, gross_week_hours - holiday_week_deduction)
        w_stat.balance_hours = w_stat.logged_hours - w_stat.target_hours
        w_stat.completion_percentage = (
            (w_stat.logged_hours / w_stat.target_hours * 100.0) if w_stat.target_hours > 0 else 100.0
        )

        summary.weekly_hours[w.sheet_name] = w_stat.logged_hours
        summary.weekly_stats[w.sheet_name] = w_stat

    summary.included_weeks = [s for s in workbook.sheet_names if s in included_sheets_order]

    summary.gross_base_hours = sum(stat.gross_days_in_month * 8.0 for stat in summary.weekly_stats.values())
    summary.holiday_count = sum(stat.holiday_days for stat in summary.weekly_stats.values())
    summary.holiday_deducted_hours = summary.holiday_count * 8.0
    summary.expected_target_hours = max(0.0, summary.gross_base_hours - summary.holiday_deducted_hours)
    summary.leave_hours = summary.leave_count * 8.0

    return summary


# --- Member row auto-match ------------------------------------------------------------

#: Starting-point confidence threshold above which an auto-detected member match is
#: trusted without asking the user to pick from the dropdown fallback (see
#: routers/brag_documents.py's preview endpoint). Tune against real name variants.
MEMBER_MATCH_CONFIDENCE_THRESHOLD = 0.72


def best_member_match(team_members: list[str], full_name: str) -> tuple[str | None, float]:
    """Fuzzy-matches the logged-in user's `User.full_name` against the workbook's row
    names (`StandupWorkbook.team_members`) using `difflib.SequenceMatcher`. Returns
    (best_name_or_None, ratio) — the caller decides what confidence counts as "auto-match"
    vs. "show the manual dropdown fallback" (see `MEMBER_MATCH_CONFIDENCE_THRESHOLD`)."""
    if not full_name or not team_members:
        return None, 0.0

    normalized_target = full_name.strip().lower()
    best_name: str | None = None
    best_ratio = 0.0
    for name in team_members:
        ratio = difflib.SequenceMatcher(None, name.strip().lower(), normalized_target).ratio()
        if ratio > best_ratio:
            best_ratio = ratio
            best_name = name

    return best_name, best_ratio
